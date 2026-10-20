"""Draw the blind error-analysis sample (SPEC 9). Dev only; never test.

Writes into --out: labels.xlsx (to fill in, one label per row from a dropdown), key.json
(item -> condition; do not open before labeling is finished) and schemas.md (DDL of every
database in the sample). Refuses to overwrite an existing labels.xlsx. Example:
  uv run --env-file .env python scripts/sample_errors.py \
    --a qwen15b_c1_zero_shot --b qwen15b_c3_lr0.0002_ep2_seed0
"""

from __future__ import annotations

import argparse
import json
import re
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from spider_sql.data import DataPaths, load_examples, sqlite_path
from spider_sql.error_analysis import LABELS, blind_items, failures, sample
from spider_sql.execution import execute
from spider_sql.schema import schema_ddl
from spider_sql.stats import load_rows

COLUMNS = (
    ("item_id", 8),
    ("db_id", 18),
    ("question", 40),
    ("gold_sql", 50),
    ("pred_sql", 50),
    ("pred_error", 30),
    ("gold_rows", 30),
    ("pred_rows", 30),
    ("label", 22),
    ("note", 30),
)
PREVIEW_ROWS = 3
CELL_CHARS = 400


def flat(sql: str) -> str:
    """Collapse whitespace. Spider gold uses double spaces around operators and C3 copies
    that style, which would reveal the condition to the labeler."""
    return re.sub(r"\s+", " ", sql).strip()


def preview(db: Path, query: str) -> tuple[str, str]:
    """(error, first rows) for one query on one database."""
    outcome = execute(db, query)
    if not outcome.ok:
        return outcome.error, ""
    rows = outcome.rows
    if not rows:
        return "", "(0 rows)"
    text = "; ".join(repr(r) for r in rows[:PREVIEW_ROWS])
    if len(rows) > PREVIEW_ROWS:
        text += f"; ... ({len(rows)} rows)"
    return "", text[:CELL_CHARS]


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--a", required=True, help="first condition run, e.g. C1")
    parser.add_argument("--b", required=True, help="second condition run, e.g. C3 seed 0")
    parser.add_argument("--n", type=int, default=100)
    parser.add_argument("--seed", type=int, default=1234)
    parser.add_argument("--results", type=Path, default=Path("results"))
    parser.add_argument("--out", type=Path, default=Path("results/error_analysis"))
    args = parser.parse_args(argv)

    out_xlsx = args.out / "labels.xlsx"
    if out_xlsx.exists():
        raise SystemExit(f"{out_xlsx} exists; refusing to overwrite labels")
    paths = DataPaths.from_env()
    examples = load_examples(paths.spider, "dev")
    preds: dict[str, dict[int, str]] = {}
    samples: dict[str, list[int]] = {}
    meta: dict[str, object] = {
        "split": "dev",
        "failure": "EX = 0 on the example's own database",
        "n_per_condition": args.n,
        "seed": args.seed,
        "conditions": {},
    }
    conditions = meta["conditions"]
    assert isinstance(conditions, dict)
    for run in (args.a, args.b):
        rows = list(load_rows(args.results / run / "dev.jsonl").values())
        fails = failures(rows)
        samples[run] = sample(fails, args.n, args.seed)
        preds[run] = {int(str(r["index"])): str(r["pred"]) for r in rows}
        conditions[run] = {"n_dev": len(rows), "n_failures": len(fails)}

    items = blind_items(samples, args.seed)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "labels"
    ws.append([name for name, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    db_ids: set[str] = set()
    for item in items:
        ex = examples[item.index]
        db = sqlite_path(paths.spider, "dev", ex.db_id)
        pred = preds[item.condition][item.index]
        _, gold_rows = preview(db, ex.query)
        pred_error, pred_rows = preview(db, pred) if pred.strip() else ("empty", "")
        ws.append(
            [
                item.item_id,
                ex.db_id,
                ex.question,
                flat(ex.query),
                flat(pred),
                pred_error,
                gold_rows,
                pred_rows,
            ]
        )
        db_ids.add(ex.db_id)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    label_col = chr(64 + [n for n, _ in COLUMNS].index("label") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    dv.add(f"{label_col}2:{label_col}{len(items) + 1}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "C2"

    args.out.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    meta["key"] = {i.item_id: {"condition": i.condition, "index": i.index} for i in items}
    (args.out / "key.json").write_text(json.dumps(meta, indent=2) + "\n", encoding="utf-8")
    ddl = [
        f"## {d}\n\n```sql\n{schema_ddl(sqlite_path(paths.spider, 'dev', d))}\n```\n"
        for d in sorted(db_ids)
    ]
    (args.out / "schemas.md").write_text(
        "# Schemas of the sampled dev databases\n\n" + "\n".join(ddl), encoding="utf-8"
    )
    print(f"{len(items)} items -> {out_xlsx}")
    for run, c in conditions.items():
        print(f"  {run}: {c['n_failures']} failures of {c['n_dev']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
