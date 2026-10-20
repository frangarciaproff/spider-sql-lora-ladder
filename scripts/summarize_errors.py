"""Summarize the manual error labels (SPEC 9) into results/error_analysis/summary.json.

Run only after every row of labels.xlsx has a label. Example:
  uv run python scripts/summarize_errors.py
"""

from __future__ import annotations

import argparse
import json
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook

from spider_sql.error_analysis import LABELS, summarize


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dir", type=Path, default=Path("results/error_analysis"))
    args = parser.parse_args(argv)

    ws = load_workbook(args.dir / "labels.xlsx", read_only=True)["labels"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    i_id, i_label = header.index("item_id"), header.index("label")
    labels = {str(r[i_id]): str(r[i_label] or "").strip() for r in rows if r[i_id]}
    meta = json.loads((args.dir / "key.json").read_text(encoding="utf-8"))
    key = {item: v["condition"] for item, v in meta["key"].items()}
    rate = {c: v["n_failures"] / v["n_dev"] for c, v in meta["conditions"].items()}
    result = summarize(labels, key, rate)

    out = {"failure": meta["failure"], "seed": meta["seed"], "labels": LABELS, **result}
    (args.dir / "summary.json").write_text(json.dumps(out, indent=2) + "\n", encoding="utf-8")
    conds = sorted(result)
    print(f"{'label':22s}" + "".join(f"{c[:28]:>30s}" for c in conds))
    for label in LABELS:
        cells = []
        for c in conds:
            entry = result[c]
            assert isinstance(entry, dict)
            v = entry["labels"][label]
            lo, hi = v["ci95"]
            cells.append(f"{v['count']:3d} ({100 * lo:4.1f}-{100 * hi:4.1f}%)")
        print(f"{label:22s}" + "".join(f"{x:>30s}" for x in cells))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
